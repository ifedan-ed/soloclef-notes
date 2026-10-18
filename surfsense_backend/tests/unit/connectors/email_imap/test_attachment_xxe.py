"""XXE on .docx/.pptx attachments: confirm it is NOT exploitable here.

``extract_attachment_text`` parses .docx/.pptx via the app's existing
``unstructured``-backed parser, which in turn uses ``python-docx``/
``python-pptx``. Both of those configure their lxml parser with
``resolve_entities=False`` (``docx/oxml/parser.py``, ``pptx/oxml/
__init__.py``) and leave ``load_dtd`` at its default (``False``), so an
external entity in a crafted document is never fetched — not from the local
filesystem, not over the network.

These tests build a minimal, real .docx/.pptx (via the normal python-docx/
python-pptx APIs) and then splice a hand-crafted ``document.xml``/
``slide1.xml`` containing a ``file://`` and an ``http://`` external entity
into the zip, so the file is read back through the actual parsing path this
connector uses — not just python-docx in isolation.
"""

from __future__ import annotations

import io
import threading
import zipfile
from email.message import EmailMessage
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

from app.connectors.email_imap import attachments

pytestmark = pytest.mark.unit

_DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


class _RecordingHTTPServer:
    """A throwaway local HTTP server that records every path it was asked
    for, so a test can assert an XXE entity was never actually fetched."""

    def __enter__(self):
        self.hits: list[str] = []
        hits = self.hits

        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                hits.append(self.path)
                self.send_response(200)
                self.end_headers()
                self.wfile.write(b"SECRET_FROM_LOCAL_SERVER")

            def log_message(self, *args):  # silence test output
                pass

        self._server = HTTPServer(("127.0.0.1", 0), Handler)
        self.port = self._server.server_address[1]
        self._thread = threading.Thread(target=self._server.serve_forever, daemon=True)
        self._thread.start()
        return self

    def __exit__(self, *exc):
        self._server.shutdown()
        self._thread.join(timeout=2)


def _docx_with_malicious_document_xml(entity_url: str) -> bytes:
    import docx

    buf = io.BytesIO()
    doc = docx.Document()
    doc.add_paragraph("placeholder")
    doc.save(buf)
    buf.seek(0)

    malicious_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<!DOCTYPE root [
  <!ENTITY xxe SYSTEM "{entity_url}">
]>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>BEFORE &xxe; AFTER</w:t></w:r></w:p>
  </w:body>
</w:document>""".encode()

    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = malicious_xml
            zout.writestr(item, data)
    return out.getvalue()


def _pptx_with_malicious_slide_xml(entity_url: str) -> bytes:
    from pptx import Presentation

    buf = io.BytesIO()
    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[1])
    slide.shapes.title.text = "placeholder"
    pres.save(buf)
    buf.seek(0)

    out = io.BytesIO()
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slides/slide1.xml":
                text = data.decode("utf-8")
                # Splice a DOCTYPE with an external entity in front of the
                # existing root element, and reference it from the title run.
                doctype = (
                    f'<!DOCTYPE root [<!ENTITY xxe SYSTEM "{entity_url}">]>'
                )
                text = text.replace(
                    "<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>",
                    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>{doctype}',
                    1,
                )
                text = text.replace("placeholder", "BEFORE &xxe; AFTER", 1)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    return out.getvalue()


async def _extract(filename: str, data: bytes, content_type: str) -> str | None:
    msg = EmailMessage()
    msg["From"] = "a@example.com"
    msg["Subject"] = "XXE probe"
    msg.set_content("See attached.")
    maintype, _, subtype = content_type.partition("/")
    msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return await attachments.extract_attachment_text(msg, filename)


async def test_docx_file_uri_external_entity_is_not_resolved():
    payload = _docx_with_malicious_document_xml("file:///etc/hostname")
    out = await _extract("probe.docx", payload, _DOCX_CONTENT_TYPE)
    # /etc/hostname's content must never appear in the extracted text.
    with open("/etc/hostname", encoding="utf-8", errors="replace") as f:
        hostname = f.read().strip()
    if hostname:
        assert hostname not in (out or "")


async def test_docx_http_external_entity_triggers_no_network_request():
    with _RecordingHTTPServer() as server:
        entity_url = f"http://127.0.0.1:{server.port}/leak"
        payload = _docx_with_malicious_document_xml(entity_url)
        out = await _extract("probe.docx", payload, _DOCX_CONTENT_TYPE)
        assert "SECRET_FROM_LOCAL_SERVER" not in (out or "")
    assert server.hits == [], (
        f"the external entity URL was actually fetched: {server.hits}"
    )


async def test_pptx_http_external_entity_triggers_no_network_request():
    with _RecordingHTTPServer() as server:
        entity_url = f"http://127.0.0.1:{server.port}/leak"
        payload = _pptx_with_malicious_slide_xml(entity_url)
        out = await _extract("probe.pptx", payload, _PPTX_CONTENT_TYPE)
        assert "SECRET_FROM_LOCAL_SERVER" not in (out or "")
    assert server.hits == [], (
        f"the external entity URL was actually fetched: {server.hits}"
    )
