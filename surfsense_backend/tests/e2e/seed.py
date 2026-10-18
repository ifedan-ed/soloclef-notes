"""Test-only seeding for the web Playwright suite: finished Studio outputs.

Generating a Studio output for real needs the sandbox and paid models, so the
suite seeds finished ones instead and tests the viewers against them. Seeding
goes through the production save path (``app.artifacts.service.save_artifact``:
the document row, the artifact, the stored files, indexing with the fake
embeddings) and then links the result the way Studio does — a ``podcast_runs``
row for audio, a ``deliverable_jobs`` row for everything else — so the Studio
list derives each status exactly as it does live.

Mounted only by ``tests/e2e/run_backend.py``; guarded by the same shared secret
as the token mint. ``tests/`` never ships in the production image.
"""

from __future__ import annotations

import io
import json
import logging
import subprocess
import tempfile
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, FastAPI, Header, HTTPException
from pydantic import BaseModel
from sqlalchemy import select

from tests.e2e.auth_mint import _expected_secret

_logger = logging.getLogger("surfsense.e2e.seed")

router = APIRouter(prefix="/__e2e__/seed", tags=["__e2e__"])

Kind = Literal[
    "audio",
    "video",
    "report",
    "report_html",
    "report_docx",
    "slides",
    "mindmap",
    "flashcards",
    "quiz",
    "data_table",
    "infographic",
]
Status = Literal["ready", "running", "failed"]


class SeedOutputRequest(BaseModel):
    workspace_id: int
    notebook_id: int
    kind: Kind
    status: Status = "ready"
    title: str | None = None
    # A report's Markdown (what the reading view shows); the default is a stub.
    markdown: str | None = None
    # An interactive report's HTML page; the default is a one-heading stub.
    html: str | None = None


class SeedOutputResponse(BaseModel):
    output_id: int
    artifact_id: int | None


# --- deterministic file bodies ---------------------------------------------------

MINDMAP_MD = """# E2E mind map
## Photosynthesis
- Light reactions
- Calvin cycle
## Respiration
- Glycolysis
- Krebs cycle
"""


def _flashcards_json(title: str) -> bytes:
    cards = [
        {"front_text": f"E2E question {i}", "back_text": f"E2E answer {i}"}
        for i in range(1, 7)
    ]
    cards[0]["hint_text"] = "E2E hint"
    return json.dumps({"schema_version": 1, "title": title, "cards": cards}).encode()


def _quiz_json(title: str) -> bytes:
    questions = [
        {
            "question_text": f"E2E quiz question {i}?",
            "options": [f"Option {i}A", f"Option {i}B", f"Option {i}C", f"Option {i}D"],
            "correct_option_index": 1,
            "explanation_text": f"Option {i}B is right because the seed says so.",
        }
        for i in range(1, 6)
    ]
    # Questions 1-3 carry per-option rationales; 4-5 are old-style (one explanation).
    for i, question in enumerate(questions[:3], start=1):
        question["option_rationales"] = [
            f"Rationale {i}A: not this one.",
            f"Rationale {i}B: this is the one.",
            f"Rationale {i}C: not this one either.",
            f"Rationale {i}D: a common mix-up.",
        ]
    return json.dumps(
        {"schema_version": 1, "title": title, "questions": questions}
    ).encode()


def _pdf(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for n in range(2):
        page = doc.new_page()
        page.insert_text((72, 72), f"{text} — page {n + 1}", fontsize=18)
    return doc.tobytes()


def _png(width: int = 800, height: int = 1200) -> bytes:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), (240, 244, 250))
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, width - 40, 200), fill=(40, 90, 160))
    draw.text((60, 100), "E2E infographic", fill=(255, 255, 255))
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Findings"
    ws.append(["Study", "Participants", "Effect"])
    ws.append(["E2E trial A", 120, 0.4])
    ws.append(["E2E trial B", 80, 0.25])
    second = wb.create_sheet("Notes")
    second.append(["Note"])
    second.append(["E2E second sheet"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("E2E report", 0)
    doc.add_paragraph("E2E Word report body.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "E2E slide deck"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ffmpeg(args: list[str], suffix: str) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / f"out{suffix}"
        subprocess.run(
            ["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", *args, str(out)],
            check=True,
            timeout=60,
        )
        return out.read_bytes()


def _mp3() -> bytes:
    return _ffmpeg(
        ["-f", "lavfi", "-i", "sine=frequency=440:duration=3", "-ac", "1", "-b:a", "32k"],
        ".mp3",
    )


def _mp4() -> bytes:
    return _ffmpeg(
        [
            "-f", "lavfi", "-i", "color=c=navy:s=320x180:d=2",
            "-f", "lavfi", "-i", "sine=frequency=330:duration=2",
            "-shortest", "-pix_fmt", "yuv420p", "-c:v", "libx264", "-c:a", "aac",
        ],
        ".mp4",
    )


def _files_for(kind: str, title: str):
    """(studio kind, artifact format, markdown, [ArtifactFileInput])."""
    from app.artifacts.service import ArtifactFileInput as F

    pdf_mime = "application/pdf"
    if kind == "audio":
        md = f"# {title}\n\n**Speaker 0:** Welcome to the E2E overview.\n\n**Speaker 1:** Thanks, glad to be here."
        return "audio", "podcast", md, [F(_mp3(), "overview.mp3", "audio/mpeg")]
    if kind == "video":
        return "video", "video", f"# {title}\n\nE2E video.", [F(_mp4(), "overview.mp4", "video/mp4")]
    if kind == "report":
        return "report", "pdf", f"# {title}\n\nE2E report.", [F(_pdf("E2E PDF report"), "report.pdf", pdf_mime)]
    if kind == "report_docx":
        return "report", "docx", f"# {title}\n\nE2E Word report.", [
            F(_docx(), "report.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            F(_pdf("E2E Word report preview"), "report.pdf", pdf_mime, role="preview"),
        ]
    if kind == "report_html":
        html = "<!doctype html><html><body><h1>E2E interactive report</h1><p>Seeded.</p></body></html>"
        return "report", "html", f"# {title}\n\nE2E interactive report.", [F(html.encode(), "report.html", "text/html")]
    if kind == "slides":
        return "slides", "pptx", f"# {title}\n\nE2E slides.", [
            F(_pptx(), "deck.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
            F(_pdf("E2E slide preview"), "deck.pdf", pdf_mime, role="preview"),
        ]
    if kind == "mindmap":
        return "mindmap", "mindmap", MINDMAP_MD, [F(_png(600, 400), "mindmap.png", "image/png")]
    if kind == "flashcards":
        return "flashcards", "flashcards", f"# {title}\n\nE2E flashcards.", [
            F(_flashcards_json(title), "flashcards.json", "application/json")
        ]
    if kind == "quiz":
        return "quiz", "quiz", f"# {title}\n\nE2E quiz.", [F(_quiz_json(title), "quiz.json", "application/json")]
    if kind == "data_table":
        return "data_table", "xlsx", f"# {title}\n\nE2E data table.", [
            F(_xlsx(), "table.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        ]
    if kind == "infographic":
        return "infographic", "infographic", f"# {title}\n\nE2E infographic.", [F(_png(), "infographic.png", "image/png")]
    raise HTTPException(status_code=400, detail=f"unknown kind {kind!r}")


@router.post("/studio-output", response_model=SeedOutputResponse)
async def seed_studio_output(
    body: SeedOutputRequest,
    x_e2e_mint_secret: str = Header(..., alias="X-E2E-Mint-Secret"),
) -> SeedOutputResponse:
    if x_e2e_mint_secret != _expected_secret():
        raise HTTPException(status_code=403, detail="invalid e2e secret")

    from app.artifacts.service import save_artifact
    from app.db import (
        DeliverableJob,
        DeliverableJobStatus,
        Workspace,
        async_session_maker,
    )
    from app.notebooks.persistence import Notebook, NotebookOutput, NotebookSource
    from app.podcasts.persistence import Podcast, PodcastStatus

    title = body.title or f"E2E {body.kind.replace('_', ' ')}"
    studio_kind, artifact_format, markdown, files = _files_for(body.kind, title)
    if body.markdown and studio_kind == "report":
        markdown = body.markdown
    if body.html and body.kind == "report_html":
        from app.artifacts.service import ArtifactFileInput

        files = [ArtifactFileInput(body.html.encode(), "report.html", "text/html")]

    async with async_session_maker() as session:
        notebook = await session.scalar(
            select(Notebook).where(
                Notebook.id == body.notebook_id,
                Notebook.workspace_id == body.workspace_id,
            )
        )
        if notebook is None:
            raise HTTPException(status_code=404, detail="notebook not found")
        owner_id = await session.scalar(
            select(Workspace.user_id).where(Workspace.id == body.workspace_id)
        )
        source_ids = list(
            (
                await session.scalars(
                    select(NotebookSource.document_id).where(
                        NotebookSource.notebook_id == notebook.id,
                        NotebookSource.document_id.is_not(None),
                    )
                )
            ).all()
        )

        artifact_id: int | None = None
        if body.status == "ready":
            saved = await save_artifact(
                session,
                workspace_id=body.workspace_id,
                thread_id=None,
                tool_call_id=f"e2e-seed-{uuid.uuid4().hex}",
                title=title,
                markdown_representation=markdown,
                files=files,
                format=artifact_format,
                extra_metadata={"duration_seconds": 3} if studio_kind == "audio" else None,
            )
            artifact_id = saved.artifact_id

        now = datetime.now(UTC)
        podcast_id = job_id = None
        if studio_kind == "audio":
            podcast = Podcast(
                title=title,
                status={
                    "ready": PodcastStatus.READY,
                    "running": PodcastStatus.RENDERING,
                    "failed": PodcastStatus.FAILED,
                }[body.status],
                workspace_id=body.workspace_id,
                artifact_id=artifact_id,
                duration_seconds=3 if body.status == "ready" else None,
                error="E2E seeded failure" if body.status == "failed" else None,
            )
            session.add(podcast)
            await session.flush()
            podcast_id = podcast.id
        else:
            job = DeliverableJob(
                kind="video" if studio_kind == "video" else "artifact",
                title=title,
                workspace_id=body.workspace_id,
                created_by_id=owner_id,
                tool_call_id=f"notebook-studio:{notebook.id}:e2e-{uuid.uuid4().hex}",
                request={"kind": studio_kind},
                status={
                    "ready": DeliverableJobStatus.READY,
                    "running": DeliverableJobStatus.RUNNING,
                    "failed": DeliverableJobStatus.FAILED,
                }[body.status],
                # A running job with no heartbeat is never reaped as an orphan
                # and never republished, so it stays "running" for the test.
                phase="rendering" if body.status == "running" else None,
                progress=40 if body.status == "running" else (100 if body.status == "ready" else 0),
                artifact_id=artifact_id,
                failure_code="generation_failed" if body.status == "failed" else None,
                finished_at=now if body.status != "running" else None,
            )
            session.add(job)
            await session.flush()
            job_id = job.id

        options = {
            "audio": {"format": "deep_dive", "length": "default"},
            "video": {"format": "explainer"},
            "report": {"format": "interactive" if body.kind == "report_html" else "document"},
            "slides": {"format": "detailed", "length": "default"},
            "infographic": {"style": "auto"},
            "quiz": {"count": "fewer", "focus": "E2E focus: the seeded questions"},
        }.get(studio_kind, {})
        output = NotebookOutput(
            notebook_id=notebook.id,
            workspace_id=body.workspace_id,
            kind=studio_kind,
            title=title,
            options=options,
            source_document_ids=source_ids,
            podcast_id=podcast_id,
            deliverable_job_id=job_id,
            created_by_id=owner_id,
            generated_by={"model": "e2e-fake-model"} if body.status == "ready" else None,
        )
        session.add(output)
        await session.commit()
        return SeedOutputResponse(output_id=output.id, artifact_id=artifact_id)


class SeedSourcePdfRequest(BaseModel):
    workspace_id: int
    notebook_id: int
    title: str = "E2E guideline"


class SeedSourcePdfResponse(BaseModel):
    document_id: int
    chunk_id: int


@router.post("/source-pdf", response_model=SeedSourcePdfResponse)
async def seed_source_pdf(
    body: SeedSourcePdfRequest,
    x_e2e_mint_secret: str = Header(..., alias="X-E2E-Mint-Secret"),
) -> SeedSourcePdfResponse:
    """A PDF source with its original file stored through the production path
    (``store_document_file``) and one chunk read from page 2, linked to the
    notebook, so a citation can open the original at that page."""
    if x_e2e_mint_secret != _expected_secret():
        raise HTTPException(status_code=403, detail="invalid e2e secret")

    from app.db import Chunk, Document, DocumentType, Workspace, async_session_maker
    from app.file_storage.service import store_document_file
    from app.notebooks.persistence import Notebook, NotebookSource

    text = f"{body.title}: use a spacer with the inhaler."
    async with async_session_maker() as session:
        notebook = await session.scalar(
            select(Notebook).where(
                Notebook.id == body.notebook_id,
                Notebook.workspace_id == body.workspace_id,
            )
        )
        if notebook is None:
            raise HTTPException(status_code=404, detail="notebook not found")
        owner_id = await session.scalar(
            select(Workspace.user_id).where(Workspace.id == body.workspace_id)
        )
        document = Document(
            title=f"{body.title}.pdf",
            document_type=DocumentType.FILE,
            content=text,
            source_markdown=text,
            content_hash=uuid.uuid4().hex,
            workspace_id=body.workspace_id,
            created_by_id=owner_id,
            status={"state": "ready"},
        )
        session.add(document)
        await session.flush()
        await store_document_file(
            session,
            document_id=document.id,
            workspace_id=body.workspace_id,
            data=_pdf(body.title),
            filename=f"{body.title}.pdf",
            mime_type="application/pdf",
            created_by_id=owner_id,
        )
        chunk = Chunk(
            content=text, document_id=document.id, position=0, page_start=2, page_end=2
        )
        session.add(chunk)
        session.add(NotebookSource(notebook_id=notebook.id, document_id=document.id))
        await session.commit()
        return SeedSourcePdfResponse(document_id=document.id, chunk_id=chunk.id)


class CompleteOutputRequest(BaseModel):
    workspace_id: int
    notebook_id: int
    output_id: int


@router.post("/complete-output", response_model=SeedOutputResponse)
async def complete_studio_output(
    body: CompleteOutputRequest,
    x_e2e_mint_secret: str = Header(..., alias="X-E2E-Mint-Secret"),
) -> SeedOutputResponse:
    """Finish a queued Studio generation (one the app started) with seeded files.

    The e2e stack cannot run the generation itself, so this plays the worker:
    it saves the kind's seeded artifact through the production save path and
    marks the output's deliverable job ready, as a finished run would.
    """
    if x_e2e_mint_secret != _expected_secret():
        raise HTTPException(status_code=403, detail="invalid e2e secret")

    from app.artifacts.service import save_artifact
    from app.db import DeliverableJob, DeliverableJobStatus, async_session_maker
    from app.notebooks.persistence import NotebookOutput

    async with async_session_maker() as session:
        output = await session.scalar(
            select(NotebookOutput).where(
                NotebookOutput.id == body.output_id,
                NotebookOutput.notebook_id == body.notebook_id,
                NotebookOutput.workspace_id == body.workspace_id,
            )
        )
        if output is None or output.deliverable_job_id is None:
            raise HTTPException(status_code=404, detail="no queued output")
        job = await session.get(DeliverableJob, output.deliverable_job_id)
        _, artifact_format, markdown, files = _files_for(output.kind, output.title)
        saved = await save_artifact(
            session,
            workspace_id=body.workspace_id,
            thread_id=None,
            tool_call_id=f"e2e-complete-{uuid.uuid4().hex}",
            title=output.title,
            markdown_representation=markdown,
            files=files,
            format=artifact_format,
        )
        job.status = DeliverableJobStatus.READY
        job.progress = 100
        job.artifact_id = saved.artifact_id
        job.finished_at = datetime.now(UTC)
        await session.commit()
        return SeedOutputResponse(output_id=output.id, artifact_id=saved.artifact_id)


def _let_report_suggestions_start():
    """Let the report view's Add start a generation on the e2e stack.

    Only the report route's sandbox check and its own dispatcher seam are
    replaced (the Studio dialog keeps its 503 path): the generation is queued
    for real but never dispatched, and ``/complete-output`` finishes it.
    """
    from unittest.mock import patch

    from app.notebooks.routes import report_view
    from app.notebooks.studio.start import Dispatchers

    patch.object(report_view, "is_sandbox_enabled", lambda: True).start()
    report_view_dispatchers = Dispatchers(
        deliverable=lambda **_kwargs: None,
        podcast_draft=lambda *_args: None,
    )
    return report_view_dispatchers


def install(app: FastAPI) -> None:
    from app.notebooks.routes.report_view import get_report_dispatchers

    app.include_router(router)
    dispatchers = _let_report_suggestions_start()
    app.dependency_overrides[get_report_dispatchers] = lambda: dispatchers
    _logger.warning(
        "[e2e] mounted POST /__e2e__/seed/studio-output and /complete-output "
        "(test-only seeding); report suggestions queue without dispatching"
    )
