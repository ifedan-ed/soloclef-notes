"""Gather the notebook's finished Studio outputs an interactive report can embed.

The newest ready output of each kind is used. Each is read on its own and a
missing, unreadable or oversized one is simply left out, so a report never
fails because an older output did. Study content (mind map, flashcards, quiz,
table) is carried as data; audio and video only as references to their stored
files, which the viewer plays through signed links rather than inlining.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.artifacts.persistence import Artifact, ArtifactFile, ArtifactFileRole
from app.artifacts.storage import open_artifact_file_stream
from app.artifacts.verification.formats.flashcards import parse_flashcards_deck
from app.artifacts.verification.formats.quiz import parse_quiz

from ...persistence import Notebook
from ..listing import list_outputs

logger = logging.getLogger(__name__)

EMBEDDABLE_KINDS = ("mindmap", "flashcards", "quiz", "data_table", "audio", "video")
_MAX_READ_BYTES = 4 * 1024 * 1024
TABLE_MAX_ROWS = 200
TABLE_MAX_COLUMNS = 12
_CELL_MAX_CHARS = 300


@dataclass(frozen=True)
class MediaEmbed:
    title: str
    artifact_id: int
    file_id: int
    mime_type: str
    duration_seconds: int | None


@dataclass(frozen=True)
class StudyEmbed:
    title: str
    # Mind map: its Markdown outline. Flashcards: (front, back, hint) triples.
    # Quiz: (question, options, correct index, explanation). Table: rows.
    payload: object


@dataclass
class StudioEmbeds:
    mindmap: StudyEmbed | None = None
    flashcards: StudyEmbed | None = None
    quiz: StudyEmbed | None = None
    data_table: StudyEmbed | None = None
    audio: MediaEmbed | None = None
    video: MediaEmbed | None = None
    kinds: list[str] = field(default_factory=list)

    @property
    def empty(self) -> bool:
        return not self.kinds


def _primary(artifact: Artifact) -> ArtifactFile | None:
    return next(
        (file for file in artifact.files if file.role is ArtifactFileRole.PRIMARY),
        None,
    )


async def _read(record: ArtifactFile) -> bytes:
    if record.size_bytes > _MAX_READ_BYTES:
        raise ValueError("file is too large to embed")
    data = bytearray()
    async for chunk in open_artifact_file_stream(record):
        data.extend(chunk)
        if len(data) > _MAX_READ_BYTES:
            raise ValueError("file is too large to embed")
    return bytes(data)


def _cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text[:_CELL_MAX_CHARS]


def read_table(data: bytes) -> list[list[str]]:
    """The first sheet's used range, bounded, as text; blank rows dropped."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows: list[list[str]] = []
        for row in sheet.iter_rows(max_col=TABLE_MAX_COLUMNS, values_only=True):
            cells = [_cell(value) for value in row]
            if any(cell.strip() for cell in cells):
                rows.append(cells)
            if len(rows) > TABLE_MAX_ROWS:
                break
    finally:
        workbook.close()
    width = max((len(row) for row in rows), default=0)
    while width and all(
        not row[width - 1].strip() for row in rows if len(row) >= width
    ):
        width -= 1
    return [(row + [""] * width)[:width] for row in rows]


async def _study(kind: str, artifact: Artifact, title: str) -> StudyEmbed | None:
    if kind == "mindmap":
        markdown = (artifact.document.source_markdown or "").strip()
        has_branches = re.search(r"^\s*[-*+] \S", markdown, re.MULTILINE)
        return StudyEmbed(title, markdown) if has_branches else None
    primary = _primary(artifact)
    if primary is None:
        return None
    data = await _read(primary)
    if kind == "flashcards":
        deck = parse_flashcards_deck(data)
        return StudyEmbed(
            title,
            [(card.front_text, card.back_text, card.hint_text) for card in deck.cards],
        )
    if kind == "quiz":
        quiz = parse_quiz(data)
        return StudyEmbed(
            title,
            [
                (
                    question.question_text,
                    list(question.options),
                    question.correct_option_index,
                    question.explanation_text,
                )
                for question in quiz.questions
            ],
        )
    if kind == "data_table":
        rows = read_table(data)
        return StudyEmbed(title, rows) if rows else None
    return None


async def collect_embeds(session: AsyncSession, notebook: Notebook) -> StudioEmbeds:
    newest: dict[str, tuple[int, str, int | None]] = {}
    for output in await list_outputs(session, notebook):
        if (
            output.kind in EMBEDDABLE_KINDS
            and output.status == "ready"
            and output.artifact_id
            and output.kind not in newest
        ):
            newest[output.kind] = (
                output.artifact_id,
                output.title,
                output.duration_seconds,
            )

    embeds = StudioEmbeds()
    if not newest:
        return embeds
    artifacts = {
        artifact.id: artifact
        for artifact in (
            await session.scalars(
                select(Artifact)
                .options(selectinload(Artifact.files), selectinload(Artifact.document))
                .where(
                    Artifact.id.in_([value[0] for value in newest.values()]),
                    Artifact.workspace_id == notebook.workspace_id,
                )
            )
        ).all()
    }
    for kind in EMBEDDABLE_KINDS:
        if kind not in newest:
            continue
        artifact_id, title, duration = newest[kind]
        artifact = artifacts.get(artifact_id)
        if artifact is None:
            continue
        try:
            if kind in ("audio", "video"):
                primary = _primary(artifact)
                if primary is None or not (primary.mime_type or "").startswith(
                    f"{kind}/"
                ):
                    continue
                embed = MediaEmbed(
                    title=title,
                    artifact_id=artifact.id,
                    file_id=primary.id,
                    mime_type=primary.mime_type,
                    duration_seconds=duration,
                )
            else:
                embed = await _study(kind, artifact, title)
        except Exception:
            logger.warning(
                "Leaving %s artifact %s out of the report",
                kind,
                artifact_id,
                exc_info=True,
            )
            continue
        if embed is not None:
            setattr(embeds, kind, embed)
            embeds.kinds.append(kind)
    return embeds
